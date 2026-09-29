import unittest
import os,tempfile
from uuid import uuid4
from io import BytesIO
from pathlib import Path
from docx import Document
_test_directory=tempfile.TemporaryDirectory(prefix="nawi-api-test-")
os.environ["DATABASE_URL"]="sqlite:///"+_test_directory.name.replace("\\","/")+"/test.db"
from fastapi.testclient import TestClient
from app.main import app
from app.db.database import engine
from app.api import attachments,reports
attachments.STORE=Path(_test_directory.name)/"attachments";attachments.STORE.mkdir(exist_ok=True)
reports.OUT=Path(_test_directory.name)/"generated_reports";reports.OUT.mkdir(exist_ok=True)

class ApiTests(unittest.TestCase):
    def test_health_login_dashboard_and_admin_rbac(self):
        with TestClient(app) as client:
            self.assertEqual(client.get("/api/health").status_code,200)
            login=client.post("/api/auth/login",data={"username":"admin@nawi.local","password":"Admin@123"})
            self.assertEqual(login.status_code,200)
            admin={"Authorization":"Bearer "+login.json()["access_token"]}
            self.assertEqual(client.get("/api/dashboard/summary",headers=admin).status_code,200)
            tech=client.post("/api/auth/login",data={"username":"technician@nawi.local","password":"Tech@123"})
            denied=client.get("/api/users",headers={"Authorization":"Bearer "+tech.json()["access_token"]})
            self.assertEqual(denied.status_code,403)
            self.assertEqual(client.post("/api/users",headers={"Authorization":"Bearer "+tech.json()["access_token"]},json={}).status_code,403)
            self.assertEqual(client.post("/api/rules",headers={"Authorization":"Bearer "+tech.json()["access_token"]},json={}).status_code,403)
            serial="T-"+uuid4().hex[:10]
            item=client.post("/api/instruments",headers=admin,json={"manufacturer":"Test Lab","model":"Scale X","serial_number":serial,"instrument_type":"Digital scale","accuracy_class":"III","capacity_unit":"kg","max_capacity":26,"min_capacity":0.2,"verification_interval_e":0.01,"display_interval_d":0.01})
            self.assertEqual(item.status_code,201,item.text)
            self.assertEqual(item.json()["number_of_verification_intervals_n"],2600)
            session=client.post("/api/tests",headers=admin,json={"instrument_id":item.json()["id"],"test_date":"2026-09-28","laboratory":"Demo lab"})
            self.assertEqual(session.status_code,201,session.text)
            tid=session.json()["id"]
            self.assertEqual(client.get(f"/api/tests/{tid}",headers={"Authorization":"Bearer "+tech.json()["access_token"]}).status_code,403)
            self.assertEqual(client.post(f"/api/tests/{tid}/review",headers={"Authorization":"Bearer "+tech.json()["access_token"]},json={"decision":"APPROVE"}).status_code,403)
            self.assertEqual(client.get(f"/api/instruments/{item.json()['id']}",headers=admin).status_code,200)
            self.assertEqual(client.patch(f"/api/instruments/{item.json()['id']}",headers=admin,json={"manufacturer":"Updated Test Lab"}).status_code,200)
            self.assertEqual(client.patch(f"/api/instruments/{item.json()['id']}",headers=admin,json={"accuracy_class":"II"}).status_code,409)
            observation=client.post(f"/api/tests/{tid}/observations",headers=admin,json={"observations":[{"test_type":"WEIGHING","test_point":"1","applied_load":1,"indicated_value":1}]})
            self.assertEqual(observation.status_code,201,observation.text)
            evaluation=client.post(f"/api/tests/{tid}/run-compliance",headers=admin)
            self.assertEqual(evaluation.status_code,200,evaluation.text)
            self.assertEqual(evaluation.json()["overall_status"],"NOT_EVALUATED")
            self.assertEqual(len(evaluation.json()["results"]),5)
            self.assertEqual(client.get(f"/api/compliance/{tid}",headers=admin).status_code,200)
            self.assertEqual(client.post(f"/api/tests/{tid}/submit",headers=admin).status_code,409)
            configured=client.post("/api/rules",headers=admin,json={"standard_name":"OIML R-76","standard_version":"TEST-FIXTURE-ONLY","rule_code":"TEST-MPE-01","test_type":"MPE","accuracy_class":"III","applicable_range":{"max_load":None,"unit":"kg"},"limit":0.1,"unit":"kg","comparison_operator":"<=","description":"Synthetic rule for automated testing only; not an OIML value.","source_reference":"Automated test fixture; not a controlled OIML citation","active":True})
            self.assertEqual(configured.status_code,201,configured.text)
            self.assertEqual(client.post("/api/rules/demo-profile",headers=admin).status_code,409)
            self.assertEqual(client.get("/api/rules",headers=admin).status_code,200)
            evaluated=client.post(f"/api/tests/{tid}/run-compliance",headers=admin)
            self.assertEqual(evaluated.status_code,200,evaluated.text)
            mpe=next(x for x in evaluated.json()["results"] if x["test_type"]=="MPE")
            self.assertEqual(mpe["status"],"PASS")
            self.assertEqual(mpe["standard_version"],"TEST-FIXTURE-ONLY")
            result=client.get(f"/api/compliance/{tid}",headers=admin).json()["results"]
            self.assertEqual(next(x for x in result if x["test_type"]=="MPE")["standard_version"],"TEST-FIXTURE-ONLY")
            client.post(f"/api/tests/{tid}/observations",headers=admin,json={"observations":[{"test_type":"WEIGHING","test_point":"2","applied_load":1,"indicated_value":1.25}]})
            failed=client.post(f"/api/tests/{tid}/run-compliance",headers=admin).json()
            self.assertEqual(next(x for x in failed["results"] if x["test_type"]=="MPE")["status"],"FAIL")
            summary=client.get("/api/dashboard/summary",headers=admin).json()
            self.assertEqual(summary["not_evaluated"],0)
            self.assertGreaterEqual(summary["failed"],1)
            self.assertEqual(client.post(f"/api/tests/{tid}/submit",headers=admin).status_code,409)
            files={"file":("../support.pdf",b"%PDF-1.4\nfixture", "application/pdf")}
            attachment=client.post("/api/attachments",headers=admin,files=files,data={"instrument_id":str(item.json()["id"])})
            self.assertEqual(attachment.status_code,201,attachment.text)
            self.assertEqual(client.get(f"/api/attachments/{attachment.json()['id']}",headers=admin).status_code,200)
            self.assertTrue(client.get("/api/users/audit",headers=admin).json())
            new_user=client.post("/api/users",headers=admin,json={"name":"QA User","email":"qa-"+uuid4().hex[:8]+"@nawi.local","password":"Testing-Password-123","role":"APPROVING_OFFICER"})
            self.assertEqual(new_user.status_code,201,new_user.text)
            self.assertEqual(client.patch(f"/api/users/{new_user.json()['id']}",headers=admin,json={"is_active":False}).status_code,200)
            self.assertEqual(client.post("/api/auth/login",data={"username":new_user.json()["email"],"password":"Testing-Password-123"}).status_code,401)
            report=client.post(f"/api/reports/from-test/{tid}",headers=admin)
            self.assertEqual(report.status_code,201,report.text)
            rid=report.json()["id"]
            self.assertEqual(client.get(f"/api/reports/{rid}/pdf",headers=admin).status_code,200)
            docx_response=client.get(f"/api/reports/{rid}/docx",headers=admin)
            self.assertEqual(docx_response.status_code,200)
            document=Document(BytesIO(docx_response.content))
            text=" ".join(p.text for p in document.paragraphs)+" "+" ".join(c.text for table in document.tables for row in table.rows for c in row.cells)
            self.assertIn("NOT_EVALUATED",text)
            self.assertIn("Synthetic demo rules are not official OIML limits",text)
            actions={x["action"] for x in client.get("/api/users/audit",headers=admin).json()}
            self.assertTrue({"RULE_CONFIGURATION_CHANGED","INSTRUMENT_UPDATED","ATTACHMENT_UPLOADED","USER_CREATED","COMPLIANCE_EVALUATION_RUN"}.issubset(actions))

    def test_complete_demo_pass_fail_reports_and_officer_review(self):
        with TestClient(app) as client:
            admin_login=client.post("/api/auth/login",data={"username":"admin@nawi.local","password":"Admin@123"}).json()
            admin={"Authorization":"Bearer "+admin_login["access_token"]}
            tech_login=client.post("/api/auth/login",data={"username":"technician@nawi.local","password":"Tech@123"}).json()
            tech={"Authorization":"Bearer "+tech_login["access_token"]}
            officer_login=client.post("/api/auth/login",data={"username":"officer@nawi.local","password":"Officer@123"}).json()
            officer={"Authorization":"Bearer "+officer_login["access_token"]}
            profile=client.post("/api/rules/demo-profile",headers=admin)
            self.assertEqual(profile.status_code,201,profile.text)
            self.assertEqual(profile.json()["created"],5)
            self.assertIn("not official oiml values",profile.json()["notice"].lower())
            self.assertEqual(client.post("/api/rules/demo-profile",headers=tech).status_code,403)
            self.assertEqual(client.post("/api/rules/demo-profile",headers=admin).json()["created"],0)
            serial="DEMO-"+uuid4().hex[:8]
            instrument=client.post("/api/instruments",headers=admin,json={"manufacturer":"A&D Weighing","model":"AD-260","serial_number":serial,"instrument_type":"Digital Platform Weighing Instrument","accuracy_class":"III","capacity_unit":"kg","max_capacity":26,"min_capacity":0.2,"verification_interval_e":0.01,"display_interval_d":0.01})
            self.assertEqual(instrument.status_code,201,instrument.text)
            self.assertEqual(instrument.json()["number_of_verification_intervals_n"],2600)
            created=client.post("/api/tests",headers=tech,json={"instrument_id":instrument.json()["id"],"test_date":"2026-09-29","laboratory":"SIH demonstration laboratory","temperature":20,"humidity":45,"pressure":1013})
            self.assertEqual(created.status_code,201,created.text)
            self.assertEqual(created.json()["operator"],"Laboratory Technician")
            tid=created.json()["id"]
            sample=[{"test_type":"WEIGHING","test_point":f"{load} kg", "applied_load":load,"indicated_value":reading} for load,reading in [(0.2,0.2),(5,5.01),(13,13.01),(26,26.01)]]
            sample += [{"test_type":"REPEATABILITY","test_point":f"Repeat {i+1}","applied_load":13,"indicated_value":v} for i,v in enumerate([13,13.01,12.99,13,13])]
            sample += [{"test_type":"ECCENTRICITY","test_point":p,"applied_load":13,"indicated_value":v} for p,v in zip(["Center","Front-left","Front-right","Rear-left","Rear-right"],[13,13.01,12.99,13,13.01])]
            sample += [{"test_type":"TARE","test_point":"Tare check","applied_load":2,"indicated_value":2.005}]
            saved=client.post(f"/api/tests/{tid}/observations",headers=tech,json={"observations":sample})
            self.assertEqual(saved.status_code,201,saved.text)
            evaluated=client.post(f"/api/tests/{tid}/run-compliance",headers=tech)
            self.assertEqual(evaluated.status_code,200,evaluated.text)
            self.assertEqual(evaluated.json()["overall_status"],"PASS")
            self.assertEqual({r["status"] for r in evaluated.json()["results"]},{"PASS"})
            detail=client.get(f"/api/tests/{tid}",headers=tech).json()
            self.assertEqual(detail["status"],"COMPLETED")
            self.assertEqual(detail["instrument_data"]["number_of_verification_intervals_n"],2600)
            fail_sample=[dict(o) for o in sample]
            fail_sample[1]["indicated_value"]=5.2
            changed=client.put(f"/api/tests/{tid}/observations",headers=tech,json={"observations":fail_sample})
            self.assertEqual(changed.status_code,200,changed.text)
            self.assertEqual(changed.json()["results"],[])
            failed=client.post(f"/api/tests/{tid}/run-compliance",headers=tech).json()
            self.assertEqual(failed["overall_status"],"FAIL")
            self.assertEqual(next(r["status"] for r in failed["results"] if r["test_type"]=="WEIGHING"),"FAIL")
            self.assertEqual(client.post(f"/api/tests/{tid}/submit",headers=tech).status_code,409)
            fail_report=client.post(f"/api/reports/from-test/{tid}",headers=tech)
            self.assertEqual(fail_report.status_code,201,fail_report.text)
            rid=fail_report.json()["id"]
            self.assertEqual(client.get("/api/reports?status=FAIL",headers=tech).json()[0]["overall_status"],"FAIL")
            self.assertEqual(client.get(f"/api/reports/{rid}/pdf",headers=tech).content[:5],b"%PDF-")
            docx=client.get(f"/api/reports/{rid}/docx",headers=tech)
            self.assertEqual(docx.status_code,200)
            document=Document(BytesIO(docx.content))
            word_text=" ".join(p.text for p in document.paragraphs)+" "+" ".join(c.text for table in document.tables for row in table.rows for c in row.cells)
            self.assertIn("OVERALL RESULT: FAIL",word_text)
            self.assertIn("NON-AUTOMATIC WEIGHING INSTRUMENT",word_text)
            self.assertIn("SIH prototype demonstration profile",word_text)
            self.assertEqual(client.get("/api/reports?status=PENDING",headers=tech).json()[0]["id"],rid)
            self.assertEqual(client.get(f"/api/reports?q={serial}",headers=tech).status_code,200)
            self.assertEqual(client.get(f"/api/reports/{rid}",headers=tech).json()["overall_status"],"FAIL")

            reset=client.put(f"/api/tests/{tid}/observations",headers=tech,json={"observations":sample})
            self.assertEqual(reset.status_code,200)
            passing=client.post(f"/api/tests/{tid}/run-compliance",headers=tech).json()
            self.assertEqual(passing["overall_status"],"PASS")
            self.assertEqual(client.post(f"/api/tests/{tid}/submit",headers=tech).status_code,200)
            self.assertEqual(client.post(f"/api/tests/{tid}/review",headers=tech,json={"decision":"APPROVE"}).status_code,403)
            reviewed=client.post(f"/api/tests/{tid}/review",headers=officer,json={"decision":"APPROVE","reason":"Checked against the displayed prototype record."})
            self.assertEqual(reviewed.status_code,200,reviewed.text)
            self.assertEqual(reviewed.json()["status"],"APPROVED")
            self.assertEqual(client.get(f"/api/reports/{rid}",headers=officer).json()["status"],"APPROVED")
            refreshed_docx=client.get(f"/api/reports/{rid}/docx",headers=officer)
            approved_doc=Document(BytesIO(refreshed_docx.content))
            approved_text=" ".join(p.text for p in approved_doc.paragraphs)+" "+" ".join(c.text for table in approved_doc.tables for row in table.rows for c in row.cells)
            self.assertIn("Approving Officer",approved_text)
            self.assertEqual(client.get("/api/reports?status=APPROVED",headers=officer).json()[0]["status"],"APPROVED")

            second=client.post("/api/tests",headers=tech,json={"instrument_id":instrument.json()["id"],"test_date":"2026-09-29","laboratory":"SIH demonstration laboratory"}).json()
            client.post(f"/api/tests/{second['id']}/observations",headers=tech,json={"observations":sample})
            self.assertEqual(client.post(f"/api/tests/{second['id']}/run-compliance",headers=tech).json()["overall_status"],"PASS")
            client.post(f"/api/tests/{second['id']}/submit",headers=tech)
            rejected=client.post(f"/api/tests/{second['id']}/review",headers=officer,json={"decision":"REJECT","reason":"Demonstration rejection path."})
            self.assertEqual(rejected.status_code,200,rejected.text)
            self.assertEqual(rejected.json()["status"],"REJECTED")
            self.assertEqual(client.get("/api/dashboard/summary",headers=admin).status_code,200)

    @classmethod
    def tearDownClass(cls):
        engine.dispose()
        _test_directory.cleanup()

if __name__=="__main__":unittest.main()
