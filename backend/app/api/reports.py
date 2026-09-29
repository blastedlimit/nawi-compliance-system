from pathlib import Path
from uuid import uuid4

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.db.database import get_db
from app.db.models import Attachment, AuditLog, Report, TestSession, Instrument, User, RuleConfiguration
from app.api.deps import current_user, roles

router=APIRouter(prefix="/api/reports",tags=["reports"])
OUT=Path(__file__).resolve().parents[2]/"generated_reports";OUT.mkdir(exist_ok=True)
NAVY=colors.HexColor("#17324d"); BLUE=colors.HexColor("#eaf1f7"); GRID=colors.HexColor("#d8e0e8")

def evaluation(t):
    statuses=[r.status for r in t.results]
    if "FAIL" in statuses:return "FAIL"
    if len(statuses)!=5 or "NOT_EVALUATED" in statuses:return "NOT_EVALUATED"
    return "PASS" if all(s=="PASS" for s in statuses) else "NOT_EVALUATED"

def _text(value):return "—" if value is None or value=="" else str(value)

def _pdf_table(rows,widths=None):
    table=Table(rows,colWidths=widths,repeatRows=1,hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),.45,GRID),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,BLUE]),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),7),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
    return table

def make_report(t,r,attachments=(),review=None,db=None):
    inst=t.instrument;pdf=OUT/(str(r.id)+".pdf");docx=OUT/(str(r.id)+".docx")
    overall=evaluation(t);results={x.test_type:x for x in t.results}
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=17,leading=21,textColor=NAVY,alignment=TA_LEFT,spaceAfter=3))
    styles.add(ParagraphStyle(name="SectionHeading",parent=styles["Heading2"],fontName="Helvetica-Bold",fontSize=10,leading=13,textColor=NAVY,spaceBefore=12,spaceAfter=6,keepWithNext=True))
    styles.add(ParagraphStyle(name="FinePrint",parent=styles["Normal"],fontSize=8,textColor=colors.HexColor("#586b7c"),leading=11))
    fields=[("Report ID",r.report_number),("Report created",str(r.created_at)),("Session created",str(t.created_at)),("Test date",str(t.test_date)),("Laboratory",t.laboratory),("Operator",t.operator.name),("Reviewer / Approving Officer",review["name"] if review else "Pending / not recorded"),("Review date",review.get("date","Pending") if review else "Pending"),("Review status",t.status),("Review comments",review["comments"] if review else t.rejection_reason)]
    instrument=[("Manufacturer",inst.manufacturer),("Model",inst.model),("Serial number",inst.serial_number),("Instrument type",inst.instrument_type),("Accuracy class",inst.accuracy_class),("Max / Min",f"{inst.max_capacity} / {inst.min_capacity} {inst.capacity_unit}"),("Verification interval e",f"{inst.verification_interval_e} {inst.capacity_unit}"),("Display interval d",f"{inst.display_interval_d} {inst.capacity_unit}"),("Number of intervals n",inst.number_of_verification_intervals_n),("Registered / modified",f"{inst.created_at} / {inst.updated_at}")]
    environment=[("Temperature",f"{_text(t.temperature)} °C"),("Relative humidity",f"{_text(t.humidity)} %RH"),("Atmospheric pressure",f"{_text(t.pressure)} hPa")]
    rules=sorted({f"{x.standard_name} / {x.standard_version}" for x in t.results})
    obs_rows=[["Test","Point / position","Applied load","Indication","Error"]]+[[o.test_type,_text(o.test_point),f"{o.applied_load:g} {inst.capacity_unit}",f"{o.indicated_value:g} {inst.capacity_unit}",f"{o.error:+g} {inst.capacity_unit}"] for o in t.observations]
    result_rows=[["Test","Calculated value","Applicable limit","Result","Rule / version"]]
    source_rows=[["Test","Rule code / version","Configuration source"]]
    for name in ("MPE","WEIGHING","REPEATABILITY","ECCENTRICITY","TARE"):
        x=results.get(name)
        result_rows.append([name,f"{x.calculated_value} {x.calculated_unit}" if x and x.calculated_value is not None else "Not available",f"{x.limit_value} {x.limit_unit}" if x and x.limit_value is not None else "Not configured",x.status if x else "NOT_EVALUATED",f"{x.rule_code} / {x.standard_version}" if x else "MISSING_RULE"])
        rule=db.get(RuleConfiguration,x.rule_id) if db and x and x.rule_id else None
        source_rows.append([name,f"{x.rule_code} / {x.standard_version}" if x else "MISSING_RULE / UNCONFIGURED",rule.source_reference if rule else "No active rule source recorded"])
    story=[Paragraph("NON-AUTOMATIC WEIGHING INSTRUMENT",styles["ReportTitle"]),Paragraph("TEST REPORT · OIML R-76 evaluation prototype",styles["Normal"]),Spacer(1,5*mm),Paragraph(f"OVERALL RESULT: {overall}",ParagraphStyle("overall",parent=styles["Heading1"],textColor=colors.HexColor("#187348" if overall=="PASS" else "#ad3434" if overall=="FAIL" else "#8a6411"),spaceAfter=5)),Paragraph("This is a software evaluation record. It is not an official government certificate. Configured demonstration rules are synthetic and are not official OIML limits.",styles["FinePrint"])]
    def section(title,rows):
        story.append(Paragraph(title,styles["SectionHeading"]));story.append(_pdf_table([["Field","Value"]]+[[str(k),str(v)] for k,v in rows],widths=[54*mm,119*mm]))
    section("1 · REPORT INFORMATION",fields);section("2 · INSTRUMENT IDENTIFICATION",instrument);section("3 · ENVIRONMENTAL CONDITIONS",environment)
    story.extend([Paragraph("4 · TEST OBSERVATIONS",styles["SectionHeading"]),_pdf_table(obs_rows,[25*mm,36*mm,31*mm,31*mm,31*mm]),Paragraph("5 · TEST RESULTS",styles["SectionHeading"]),_pdf_table(result_rows,[27*mm,32*mm,30*mm,24*mm,42*mm])])
    story.extend([Spacer(1,3*mm),Paragraph("Evaluation configuration: "+("; ".join(rules) if rules else "unconfigured"),styles["FinePrint"]),Paragraph("6 · CONFIGURED RULE SOURCES",styles["SectionHeading"]),_pdf_table(source_rows,[24*mm,60*mm,89*mm]),Paragraph("7 · APPROVAL / REVIEW",styles["SectionHeading"]),_pdf_table([["Status","Officer","Review date","Decision comments"],[t.status,fields[6][1],fields[7][1],fields[9][1] or "—"]],[28*mm,42*mm,38*mm,47*mm]),Paragraph("8 · SUPPORTING DOCUMENTS",styles["SectionHeading"]),_pdf_table([["File","Type","Uploaded"]]+([[a.filename,a.content_type,str(a.created_at)] for a in attachments] or [["No supporting documents linked","—","—"]]),[65*mm,55*mm,35*mm])])
    def page(canvas,doc):
        canvas.saveState();w,h=A4;canvas.setStrokeColor(GRID);canvas.line(18*mm,h-16*mm,w-18*mm,h-16*mm);canvas.setFont("Helvetica-Bold",8);canvas.setFillColor(NAVY);canvas.drawString(18*mm,h-12*mm,"NAWI COMPLIANCE · TEST RECORD");canvas.setFont("Helvetica",8);canvas.setFillColor(colors.HexColor("#586b7c"));canvas.drawString(18*mm,10*mm,f"{r.report_number} · Prototype record · Not a legal certificate");canvas.drawRightString(w-18*mm,10*mm,f"Page {doc.page}");canvas.restoreState()
    SimpleDocTemplate(str(pdf),pagesize=A4,leftMargin=18*mm,rightMargin=18*mm,topMargin=22*mm,bottomMargin=17*mm,title=f"NAWI Test Report {r.report_number}",author="NAWI Compliance prototype").build(story,onFirstPage=page,onLaterPages=page)

    d=Document();sec=d.sections[0];sec.header.paragraphs[0].text="NAWI COMPLIANCE  ·  LABORATORY TEST RECORD";sec.footer.paragraphs[0].text=f"{r.report_number}  |  Prototype record · Not a legal certificate"
    sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
    normal=d.styles["Normal"];normal.font.name="Aptos";normal.font.size=Pt(9);normal.font.color.rgb=RGBColor(44,62,80)
    title=d.add_heading("NON-AUTOMATIC WEIGHING INSTRUMENT",0);title.alignment=WD_ALIGN_PARAGRAPH.LEFT
    d.add_heading("TEST REPORT · OIML R-76 evaluation prototype",2)
    d.add_heading(f"OVERALL RESULT: {overall}",1)
    d.add_paragraph("Software evaluation record only; not an official government certificate. Synthetic demo rules are not official OIML limits.")
    def word_table(heading,rows):
        d.add_heading(heading,2);tab=d.add_table(rows=0,cols=2);tab.style="Light Shading Accent 1"
        for a,b in rows:
            cells=tab.add_row().cells;cells[0].text=str(a);cells[1].text=str(b)
    word_table("1 · Report Information",fields);word_table("2 · Instrument Identification",instrument);word_table("3 · Environmental Conditions",environment)
    d.add_heading("4 · Test Observations",2);tab=d.add_table(rows=1,cols=5);tab.style="Light Shading Accent 1"
    for c,v in zip(tab.rows[0].cells,obs_rows[0]):c.text=v
    for row in obs_rows[1:]:
        cells=tab.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    d.add_heading("5 · Test Results",2);tab=d.add_table(rows=1,cols=5);tab.style="Light Shading Accent 1"
    for c,v in zip(tab.rows[0].cells,result_rows[0]):c.text=v
    for row in result_rows[1:]:
        cells=tab.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    d.add_paragraph("Evaluation configuration: "+("; ".join(rules) if rules else "unconfigured"))
    d.add_heading("6 · Configured Rule Sources",2);tab=d.add_table(rows=1,cols=3);tab.style="Light Shading Accent 1"
    for c,v in zip(tab.rows[0].cells,source_rows[0]):c.text=v
    for row in source_rows[1:]:
        cells=tab.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    word_table("7 · Approval / Review",[("Status",t.status),("Officer",fields[6][1]),("Review date",fields[7][1]),("Decision comments",fields[9][1] or "—")])
    d.add_heading("8 · Supporting Documents",2);tab=d.add_table(rows=1,cols=3);tab.style="Light Shading Accent 1"
    for c,v in zip(tab.rows[0].cells,["File","Type","Uploaded"]):c.text=v
    for row in ([[a.filename,a.content_type,str(a.created_at)] for a in attachments] or [["No supporting documents linked","—","—"]]):
        cells=tab.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    d.save(docx);r.pdf_path=str(pdf);r.docx_path=str(docx)

@router.get("")
def listing(q:str="",status:str="",db:Session=Depends(get_db),u:User=Depends(current_user)):
    query=db.query(Report).join(TestSession)
    if u.role=="LAB_TECHNICIAN":query=query.filter(TestSession.operator_id==u.id)
    if q:query=query.filter((Report.report_number.contains(q))|(TestSession.instrument.has(Instrument.serial_number.contains(q)))|(TestSession.instrument.has(Instrument.model.contains(q))))
    if status=="PENDING":query=query.filter(Report.status.notin_(["APPROVED","REJECTED"]))
    elif status in {"APPROVED","REJECTED"}:query=query.filter(Report.status==status)
    rows=[]
    for r in query.order_by(Report.created_at.desc()).all():
        t=r.test_session;rows.append({"id":r.id,"report_number":r.report_number,"status":r.status,"overall_status":evaluation(t),"test_session_id":t.id,"serial_number":t.instrument.serial_number,"model":t.instrument.model,"operator":t.operator.name,"test_date":str(t.test_date),"created_at":r.created_at})
    if status in {"PASS","FAIL"}:rows=[r for r in rows if r["overall_status"]==status]
    return rows

@router.post("/from-test/{test_id}",status_code=201)
def generate(test_id:int,db:Session=Depends(get_db),u:User=Depends(roles("ADMIN","LAB_TECHNICIAN","APPROVING_OFFICER"))):
    t=db.get(TestSession,test_id)
    if not t:raise HTTPException(404,"Test session not found")
    if u.role=="LAB_TECHNICIAN" and t.operator_id!=u.id:raise HTTPException(403,"You can only generate reports for your own tests")
    if not t.results:raise HTTPException(409,"Run compliance before generating a report")
    r=Report(test_session_id=t.id,report_number=f"NAWI-{t.test_date:%Y%m%d}-{uuid4().hex[:6].upper()}",status=t.status);db.add(r);db.flush()
    audit=db.query(AuditLog).filter(AuditLog.entity_type=="TEST",AuditLog.entity_id==str(t.id),AuditLog.action.in_(["REPORT_APPROVED","REPORT_REJECTED"])).order_by(AuditLog.timestamp.desc()).first()
    review={"name":db.get(User,audit.user_id).name if audit and audit.user_id else "Pending / not recorded","comments":audit.details if audit else "","date":str(audit.timestamp)} if audit else None
    attachments=db.query(Attachment).filter((Attachment.test_session_id==t.id)|(Attachment.instrument_id==t.instrument_id)).order_by(Attachment.created_at).all()
    make_report(t,r,attachments,review,db);db.add(AuditLog(user_id=u.id,action="REPORT_GENERATED",entity_type="REPORT",entity_id=str(r.id)));db.commit()
    return {"id":r.id,"report_number":r.report_number,"pdf_path":r.pdf_path,"docx_path":r.docx_path}

@router.get("/{report_id}")
def detail(report_id:int,db:Session=Depends(get_db),u:User=Depends(current_user)):
    r=db.get(Report,report_id)
    if not r:raise HTTPException(404,"Report not found")
    if u.role=="LAB_TECHNICIAN" and r.test_session.operator_id!=u.id:raise HTTPException(403,"You can only access your own reports")
    t=r.test_session
    return {"id":r.id,"report_number":r.report_number,"status":r.status,"overall_status":evaluation(t),"test_session_id":t.id,"test_date":str(t.test_date),"instrument":{"manufacturer":t.instrument.manufacturer,"model":t.instrument.model,"serial_number":t.instrument.serial_number},"results":[{"test_type":x.test_type,"status":x.status,"calculated_value":x.calculated_value,"limit_value":x.limit_value,"details":x.details} for x in t.results]}

@router.get("/{report_id}/{fmt}")
def download(report_id:int,fmt:str,db:Session=Depends(get_db),u:User=Depends(current_user)):
    if fmt not in ("pdf","docx"):raise HTTPException(404,"Format not found")
    r=db.get(Report,report_id)
    if not r:raise HTTPException(404,"Report not found")
    if u.role=="LAB_TECHNICIAN" and r.test_session.operator_id!=u.id:raise HTTPException(403,"You can only access your own reports")
    path=r.pdf_path if fmt=="pdf" else r.docx_path
    if not path or not Path(path).is_file():raise HTTPException(404,"Report file is not available")
    return FileResponse(path,filename=f"{r.report_number}.{fmt}")
